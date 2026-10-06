"""Excel / CSV export of leads."""

import csv
import io
from collections.abc import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.db.models import LeadRecord

# (column title, attribute on LeadRecord, column width)
COLUMNS = [
    ("ID", "id", 6),
    ("Name", "name", 22),
    ("Company", "company", 20),
    ("Category", "category", 11),
    ("Platform", "platform", 10),
    ("Fit", "fit_score", 5),
    ("Fit reason", "fit_reason", 40),
    ("Pain points", "pain_points", 40),
    ("About them", "about_them", 50),
    ("How I can help", "how_i_can_help", 50),
    ("Channel", "channel", 14),
    ("Subject", "subject", 25),
    ("Message", "message", 70),
    ("Public email", "public_email", 26),
    ("Profile URL", "profile_url", 35),
    ("Website", "company_website", 28),
    ("Signal", "signal", 50),
    ("Source URL", "source_url", 40),
    ("Status", "status", 9),
    ("Notes", "notes", 25),
    ("Found at", "created_at", 18),
]


def _value(lead: LeadRecord, attr: str):
    v = getattr(lead, attr)
    if isinstance(v, list):
        return "; ".join(map(str, v))
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d %H:%M")
    return v


def leads_to_xlsx(leads: Iterable[LeadRecord]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append([title for title, _, _ in COLUMNS])
    for lead in leads:
        ws.append([_value(lead, attr) for _, attr, _ in COLUMNS])

    # Make it readable: bold header, wrapped text, sensible widths, frozen header row.
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for i, (_, _, width) in enumerate(COLUMNS, start=1):
        letter = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[letter].width = width
        ws.cell(row=1, column=i).font = Font(bold=True)
        ws.cell(row=1, column=i).fill = header_fill
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def leads_to_csv(leads: Iterable[LeadRecord]) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([title for title, _, _ in COLUMNS])
    for lead in leads:
        writer.writerow([_value(lead, attr) for _, attr, _ in COLUMNS])
    return buf.getvalue()
